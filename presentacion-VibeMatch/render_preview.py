# -*- coding: utf-8 -*-
"""Genera una vista previa HTML del PPTX (para revisar diseño y presentar sin PowerPoint)."""
from __future__ import annotations

import html
import os
import shutil

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu

BASE = os.path.dirname(os.path.abspath(__file__))
PPTX = os.path.join(BASE, "VibeMatch-Presentacion.pptx")
OUT = os.path.join(BASE, "preview.html")
IMG_DIR = os.path.join(BASE, "preview_assets")

SCALE = 120.0  # px por pulgada -> 1600 x 900


def emu_px(v) -> float:
    return float(v) / 914400.0 * SCALE


def rgba_from_parent(parent):
    if parent is None:
        return None
    srgb = parent.find(qn("a:srgbClr"))
    if srgb is None:
        return None
    val = srgb.get("val")
    if not val or len(val) != 6:
        return None
    r, g, b = int(val[0:2], 16), int(val[2:4], 16), int(val[4:6], 16)
    alpha_el = srgb.find(qn("a:alpha"))
    alpha = int(alpha_el.get("val")) / 100000.0 if alpha_el is not None else 1.0
    return f"rgba({r},{g},{b},{alpha:.3f})"


def fill_css(shape):
    spPr = shape._element.spPr
    solid = spPr.find(qn("a:solidFill"))
    if solid is not None:
        c = rgba_from_parent(solid)
        if c:
            return c
    grad = spPr.find(qn("a:gradFill"))
    if grad is not None:
        stops = []
        gs_lst = grad.find(qn("a:gsLst"))
        if gs_lst is not None:
            for gs in gs_lst.findall(qn("a:gs")):
                c = rgba_from_parent(gs)
                if c:
                    pos = int(gs.get("pos", "0")) / 100000.0
                    stops.append(f"{c} {pos*100:.1f}%")
        if stops:
            lin = grad.find(qn("a:lin"))
            ang = int(lin.get("ang", "0")) / 60000.0 if lin is not None else 0
            css_angle = (ang + 90) % 360
            return f"linear-gradient({css_angle:.0f}deg, {', '.join(stops)})"
    return None


def line_css(shape):
    spPr = shape._element.spPr
    ln = spPr.find(qn("a:ln"))
    if ln is None:
        return None
    if ln.find(qn("a:noFill")) is not None:
        return None
    solid = ln.find(qn("a:solidFill"))
    if solid is None:
        return None
    c = rgba_from_parent(solid)
    if not c:
        return None
    w_emu = int(ln.get("w", "12700"))
    w_pt = w_emu / 12700.0
    return f"{max(0.5, w_pt * SCALE / 72.0):.2f}px solid {c}"


def radius_css(shape, w, h):
    try:
        ast = shape.auto_shape_type
    except Exception:
        return "0"
    name = str(ast)
    if "OVAL" in name:
        return "50%"
    if "ROUNDED_RECTANGLE" in name:
        try:
            adj = shape.adjustments[0]
        except Exception:
            adj = 0.06
        return f"{adj * min(w, h):.1f}px"
    return "0"


def text_html(shape, w, h):
    tf = shape.text_frame
    if not tf.text.strip():
        return ""
    ml = emu_px(tf.margin_left)
    mr = emu_px(tf.margin_right)
    mt = emu_px(tf.margin_top)
    mb = emu_px(tf.margin_bottom)
    anchor = tf.vertical_anchor
    justify = "flex-start"
    if anchor == MSO_ANCHOR.MIDDLE:
        justify = "center"
    elif anchor == MSO_ANCHOR.BOTTOM:
        justify = "flex-end"

    paras = []
    for p in tf.paragraphs:
        align = "left"
        if p.alignment is not None:
            a = str(p.alignment)
            if "CENTER" in a:
                align = "center"
            elif "RIGHT" in a:
                align = "right"
            elif "JUSTIFY" in a:
                align = "justify"
        ls = p.line_spacing
        if ls is None:
            line_height = 1.2
        elif isinstance(ls, float):
            line_height = ls * 1.2
        else:
            line_height = emu_px(ls) / max(1.0, 1.0)
            line_height = max(1.0, line_height / 16.0)
        after = emu_px(p.space_after) if p.space_after is not None else 0
        spans = []
        for r in p.runs:
            if not r.text:
                continue
            size = r.font.size.pt if r.font.size else 14
            color = "#f4f4f8"
            try:
                if r.font.color and r.font.color.rgb:
                    color = "#" + str(r.font.color.rgb)
            except Exception:
                pass
            family = r.font.name or "Segoe UI"
            if "Consolas" in family:
                family_css = "Consolas, monospace"
            elif "Semibold" in family:
                family_css = "'Segoe UI Semibold', 'Segoe UI', sans-serif"
            else:
                family_css = "'Segoe UI', Arial, sans-serif"
            weight = 700 if r.font.bold else 400
            spc = r._r.find(qn("a:rPr"))
            tracking = ""
            if spc is not None and spc.get("spc"):
                tracking = f"letter-spacing:{int(spc.get('spc'))/100*SCALE/72:.2f}px;"
            spans.append(
                f'<span style="font-family:{family_css};font-weight:{weight};'
                f'font-size:{size*SCALE/72:.2f}px;color:{color};{tracking}">'
                f"{html.escape(r.text)}</span>"
            )
        if not spans:
            continue
        paras.append(
            f'<p style="margin:0 0 {after:.2f}px 0;text-align:{align};'
            f'line-height:{line_height:.3f};">{"".join(spans)}</p>'
        )
    if not paras:
        return ""
    return (
        f'<div class="tf" style="padding:{mt:.2f}px {mr:.2f}px {mb:.2f}px {ml:.2f}px;'
        f'justify-content:{justify};">{"".join(paras)}</div>'
    )


def render():
    prs = Presentation(PPTX)
    if os.path.isdir(IMG_DIR):
        shutil.rmtree(IMG_DIR)
    os.makedirs(IMG_DIR, exist_ok=True)

    slides_html = []
    for si, slide in enumerate(prs.slides, 1):
        bg = "#07070d"
        try:
            if slide.background.fill.type is not None:
                bg = "#" + str(slide.background.fill.fore_color.rgb)
        except Exception:
            pass
        parts = []
        for i, shape in enumerate(slide.shapes):
            x, y = emu_px(shape.left), emu_px(shape.top)
            w, h = emu_px(shape.width), emu_px(shape.height)
            if shape.shape_type == 13:  # PICTURE
                try:
                    img = shape.image
                    ext = img.ext
                    name = f"s{si:02d}_{i:02d}.{ext}"
                    with open(os.path.join(IMG_DIR, name), "wb") as fh:
                        fh.write(img.blob)
                    parts.append(
                        f'<img class="shape" src="preview_assets/{name}" '
                        f'style="left:{x:.2f}px;top:{y:.2f}px;width:{w:.2f}px;height:{h:.2f}px;" />'
                    )
                except Exception as exc:
                    print("img error", si, i, exc)
                continue

            # conectores
            if shape.shape_type == 9 or "Connector" in type(shape).__name__:
                try:
                    x1, y1 = emu_px(shape.begin_x), emu_px(shape.begin_y)
                    x2, y2 = emu_px(shape.end_x), emu_px(shape.end_y)
                except Exception:
                    x1, y1, x2, y2 = x, y, x + w, y + h
                stroke = line_css(shape) or "1px solid #6D7ED3"
                color = stroke.split("solid ")[-1] if "solid " in stroke else "#6D7ED3"
                width = stroke.split("px")[0] if "px" in stroke else "1"
                parts.append(
                    f'<svg class="shape" style="left:0;top:0;width:1600px;height:900px;overflow:visible;pointer-events:none;">'
                    f'<line x1="{x1:.2f}" y1="{y1:.2f}" x2="{x2:.2f}" y2="{y2:.2f}" '
                    f'stroke="{color}" stroke-width="{width}" /></svg>'
                )
                continue

            fill = fill_css(shape)
            border = line_css(shape)
            radius = radius_css(shape, w, h)
            styles = [
                f"left:{x:.2f}px", f"top:{y:.2f}px", f"width:{w:.2f}px", f"height:{h:.2f}px",
                f"border-radius:{radius}",
            ]
            if fill:
                styles.append(f"background:{fill}")
            if border:
                styles.append(f"border:{border}")
            inner = text_html(shape, w, h)
            parts.append(f'<div class="shape" style="{";".join(styles)};">{inner}</div>')

        slides_html.append(
            f'<section class="slide" data-n="{si}" style="background:{bg};">{"".join(parts)}</section>'
        )

    doc = f"""<!doctype html>
<html lang="es">
<head>
<meta charset="utf-8" />
<meta name="viewport" content="width=device-width, initial-scale=1" />
<title>VibeMatch · Vista previa de la presentación</title>
<style>
  html, body {{ margin:0; padding:0; background:#000; }}
  body {{ font-family:'Segoe UI', Arial, sans-serif; }}
  .slide {{ position:relative; width:1600px; height:900px; overflow:hidden; display:none; }}
  .slide.active {{ display:block; }}
  .shape {{ position:absolute; box-sizing:border-box; }}
  .tf {{ position:absolute; inset:0; display:flex; flex-direction:column; box-sizing:border-box; }}
  .tf p {{ white-space:normal; }}
  .hint {{ position:fixed; left:12px; bottom:10px; color:#777; font:11px Consolas, monospace; z-index:50; }}
  @page {{ size: 1600px 900px; margin: 0; }}
  @media print {{
    html, body {{ background:#000; -webkit-print-color-adjust:exact; print-color-adjust:exact; }}
    .slide {{ display:block !important; page-break-after:always; break-after:page; margin:0; }}
    .slide:last-of-type {{ page-break-after:auto; break-after:auto; }}
    .hint {{ display:none; }}
  }}
</style>
</head>
<body>
{''.join(slides_html)}
<div class="hint">← → para navegar · F para pantalla completa</div>
<script>
  const params = new URLSearchParams(location.search);
  if (params.get('clean') === '1') {{
    const hint = document.querySelector('.hint');
    if (hint) hint.style.display = 'none';
  }}
  let current = Math.max(1, Math.min({len(slides_html)}, parseInt(params.get('slide') || '1', 10) || 1));
  const slides = [...document.querySelectorAll('.slide')];
  function show(n) {{
    current = Math.max(1, Math.min(slides.length, n));
    slides.forEach((s, i) => s.classList.toggle('active', i + 1 === current));
    try {{ history.replaceState(null, '', '?slide=' + current); }} catch (e) {{}}
  }}
  document.addEventListener('keydown', (e) => {{
    if (e.key === 'ArrowRight' || e.key === 'PageDown' || e.key === ' ') show(current + 1);
    if (e.key === 'ArrowLeft' || e.key === 'PageUp') show(current - 1);
    if (e.key.toLowerCase() === 'f') document.documentElement.requestFullscreen?.();
  }});
  show(current);
</script>
</body>
</html>
"""
    with open(OUT, "w", encoding="utf-8") as fh:
        fh.write(doc)
    print("OK:", OUT, "| slides:", len(slides_html))


if __name__ == "__main__":
    render()
